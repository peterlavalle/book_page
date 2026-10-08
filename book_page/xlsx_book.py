import contextlib
import shutil
from collections.abc import Generator as gen
from pathlib import Path

import openpyxl

from .book import Book, Page


class XLSXBook(Book):
    def __init__(self, file: str | Path, read_only=True):
        if isinstance(file, str):
            file = Path(file)
        if not file.is_file():
            raise RuntimeError(f"xlsx file not found {file}")
        self._file = file
        self._book = openpyxl.load_workbook(file, read_only=read_only)

    def keys(self):
        return [page.title for page in self._book.worksheets]

    def stream_copy(self, into: str | Path) -> gen[tuple[str, int, list[str]]]:
        if not isinstance(into, Path):
            into = Path(into)

        if not into.name.endswith(".xlsx"):
            into = into / self._file.name

        shutil.copy2(self._file, into)
        the_copy = openpyxl.load_workbook(into)

        for name in the_copy.sheetnames:
            page = the_copy[name]

            for r in range(page.max_row):
                data = [page.cell(r + 1, c + 1).value for c in range(page.max_column)]
                yield name, r, data
                for c, val in enumerate(data):
                    page.cell(r + 1, c + 1).value = val
        the_copy.save(into)

    @contextlib.contextmanager
    def __getitem__(self, name):

        if len(name) > 31:
            raise RuntimeError("excel limits you to 31 character page names")

        class XLSXPage(Page):
            def stream_rows(self) -> gen[tuple[int, list[str]]]:
                for row in range(self._page.max_row):
                    yield (
                        row,
                        [
                            self._page.cell(row + 1, col + 1).value
                            for col in range(self._page.max_column)
                        ],
                    )

        yield XLSXPage(name, self._book[name])
